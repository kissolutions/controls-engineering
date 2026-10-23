#!/usr/bin/env python3
"""Optional project PER consistency review; replaces in-workbook SB2 diagnostics.

Reports findings but does not modify files or declare engineering acceptance.
Usage: python scripts/check_project_per.py /path/to/project_PER.xlsx
"""
from collections import Counter
from pathlib import Path
import sys
from openpyxl import load_workbook

def check(path):
    w = load_workbook(path, read_only=True, data_only=False)
    for name in ("WORK PACKAGES", "WP PROGRESS"):
        if name not in w:
            raise ValueError(f"Missing register {name}")
    wp = w["WORK PACKAGES"]
    pr = w["WP PROGRESS"]
    defined = [str(wp.cell(r, 1).value).strip() for r in range(2, 101) if wp.cell(r, 1).value not in (None, "")]
    entered = [str(pr.cell(r, 1).value).strip() for r in range(2, 101) if pr.cell(r, 1).value not in (None, "")]
    expected = set(defined)
    actual = set(entered)
    findings = []
    for label, values in (("WORK PACKAGES", defined), ("WP PROGRESS", entered)):
        for value, count in Counter(values).items():
            if count > 1:
                findings.append(f"{label}: {value} occurs {count} times")
    for value in sorted(expected - actual):
        findings.append(f"No WP PROGRESS row for {value}")
    for value in sorted(actual - expected):
        findings.append(f"Orphan WP PROGRESS row: {value}")
    for r in range(2, 101):
        value = pr.cell(r, 1).value
        if value in (None, ""):
            continue
        for col in range(5, 11):
            cell = pr.cell(r, col)
            v = cell.value
            if v in (None, "") or (isinstance(v, str) and v.strip() == "--"):
                continue
            if isinstance(v, (int, float)) and not isinstance(v, bool) and 0 <= v <= 1:
                continue
            findings.append(f"{cell.coordinate}: {v!r} is not blank, -- or numeric 0..1")
    print("PROJECT PER CONSISTENCY REVIEW", Path(path).name)
    print("Work Packages:", len(defined), "| Progress rows:", len(entered))
    if not findings:
        print("No tested WP identity/progress-input issues found")
    else:
        for finding in findings:
            print(" -", finding)
    print("This diagnostic does not evaluate engineering validity, source completeness or Excel formula recalculation.")
    return int(bool(findings))

if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit("Usage: check_project_per.py PROJECT_PER.xlsx")
    sys.exit(check(sys.argv[1]))
