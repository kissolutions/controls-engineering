#!/usr/bin/env python3
"""Verify a released blank PER template, rather than a populated project copy."""
import re
import sys
import zipfile
from pathlib import Path
from openpyxl import load_workbook

SHEETS = [
    "DASHBOARD", "SCOPE DASHBOARD", "WP PROGRESS", "BLOCKED ITEMS REPORT",
    "SCOPES", "WORK PACKAGES", "SYSTEM MODEL", "RFI REGISTER", "ISSUES",
    "PUNCH LIST", "REGISTRY GUIDE", "SOURCES", "RECORDS", "LINKS",
    "CHANGE LOG", "BASELINES", "INTAKE QUEUE",
]
START = {
    "DASHBOARD": 3, "SCOPE DASHBOARD": 5, "WP PROGRESS": 2,
    "BLOCKED ITEMS REPORT": 5, "SCOPES": 5, "WORK PACKAGES": 2,
    "SYSTEM MODEL": 5, "RFI REGISTER": 5, "ISSUES": 5, "PUNCH LIST": 3,
    "SOURCES": 5, "RECORDS": 5, "LINKS": 5, "CHANGE LOG": 5,
    "BASELINES": 5, "INTAKE QUEUE": 5,
}
REMOVED = ("RFI LEGACY", "SB1 FORMULA TESTS", "SB2 CHECKS")
LEAK = re.compile(r"https?://(?:drive|docs)\.google\.com/|file:///|\\\\[^\\]", re.I)

def intentional_default(sheet, cell):
    if sheet == "PUNCH LIST":
        return ((cell.column in (20, 21, 22) and cell.value == 1)
                or (cell.column == 17 and cell.value == "Open")
                or (cell.coordinate == "B3" and cell.value == 1))
    return False

def validate(path):
    issues = []
    w = load_workbook(path)
    if w.sheetnames != SHEETS:
        issues.append("Incorrect sheet set/order; expected 17 operational worksheets")
    if w._external_links:
        issues.append("External workbook link remains")
    try:
        with zipfile.ZipFile(path) as z:
            damaged = z.testzip()
            if damaged:
                issues.append(f"Damaged ZIP member: {damaged}")
    except Exception as exc:
        issues.append(f"Unreadable XLSX ZIP: {exc}")
    formulas = 0
    for s in w:
        for row in s:
            for c in row:
                if c.value is None:
                    continue
                value = str(c.value)
                if any(tag in value for tag in REMOVED):
                    issues.append(f"{s.title}!{c.coordinate} references removed worksheet")
                if c.data_type == "f":
                    formulas += 1
                    if "#REF!" in value:
                        issues.append(f"{s.title}!{c.coordinate} broken formula")
                    continue
                if s.title in START and c.row >= START[s.title] and not intentional_default(s.title, c):
                    issues.append(f"{s.title}!{c.coordinate} nonformula project data")
                if LEAK.search(value) or c.hyperlink:
                    issues.append(f"{s.title}!{c.coordinate} leaked path or hyperlink")
                if c.comment and s.title in START:
                    issues.append(f"{s.title}!{c.coordinate} has a comment")
    for name in w.defined_names.values():
        if any(tag in str(name.attr_text) for tag in REMOVED):
            issues.append(f"Defined name {name.name} references removed sheet")
    if formulas != 14127:
        issues.append(f"Unexpected formula count {formulas}; v1.2 distribution has 14127")
    if sum(len(s.data_validations.dataValidation) for s in w) != 29:
        issues.append("Unexpected number of data validation rules")
    if w["DASHBOARD"]["U2"].value != "Engineering Scope / Scope of Work":
        issues.append("Dashboard v1.2 header mismatch")
    if w.properties.creator not in (None, "KIS Solutions") or w.properties.lastModifiedBy not in (None, "KIS Solutions"):
        issues.append("Personal editor/creator property remains")
    print(("FAIL" if issues else "PASS"), Path(path).name, len(w.sheetnames), "sheets,", formulas, "formulas")
    for issue in issues[:30]:
        print("  -", issue)
    if len(issues) > 30:
        print(f"  ... and {len(issues)-30} additional findings")
    return int(bool(issues))

if __name__ == "__main__":
    if len(sys.argv) == 2:
        source = Path(sys.argv[1])
    elif len(sys.argv) == 1:
        source = Path(__file__).resolve().parent.parent / "templates/project-engineering-register/PROJECT_PER_v1.2.xlsx"
    else:
        sys.exit("Usage: validate_blank_per.py [PROJECT_PER_v1.2.xlsx]")
    sys.exit(validate(source))
